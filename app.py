import io
import os
import re
import sqlite3
import pandas as pd
import streamlit as st

DB_FILE = "neo_master.db"
MASTER_CSV = "master_list.csv"

st.set_page_config(page_title="ID Lookup", page_icon="🔎", layout="wide")

def db():
    con = sqlite3.connect(DB_FILE)
    con.execute("""
        CREATE TABLE IF NOT EXISTS neo_master (
            neo_id TEXT PRIMARY KEY,
            name TEXT NOT NULL
        )
    """)
    con.commit()
    return con

def normalize_id(x):
    return str(x).strip().upper()

def seed_database():
    con = db()
    count = con.execute("SELECT COUNT(*) FROM neo_master").fetchone()[0]
    if count == 0 and os.path.exists(MASTER_CSV):
        df = pd.read_csv(MASTER_CSV, dtype=str).fillna("")
        rows = [(normalize_id(r["neo_id"]), str(r["name"]).strip())
                for _, r in df.iterrows()
                if str(r["neo_id"]).strip() and str(r["name"]).strip()]
        con.executemany("INSERT OR REPLACE INTO neo_master(neo_id,name) VALUES (?,?)", rows)
        con.commit()
    return con

def get_master(con):
    return pd.read_sql_query(
        "SELECT neo_id, name FROM neo_master ORDER BY neo_id", con
    )

def extract_ids_from_dataframe(df):
    # Prefer a column whose header contains "neo" and "id".
    cols = [str(c) for c in df.columns]
    preferred = [c for c in cols if "neo" in c.lower() and "id" in c.lower()]
    series = df[preferred[0]] if preferred else None

    if series is not None:
        vals = [normalize_id(x) for x in series.dropna().tolist()]
        vals = [x for x in vals if re.fullmatch(r"[A-Z0-9]{5,20}", x)]
        if vals:
            return vals

    # Fallback: scan every cell for NEO-ID-looking tokens.
    text = "\n".join(
        str(v) for v in df.astype(str).fillna("").values.flatten()
    ).upper()
    return sorted(set(re.findall(r"\b[A-Z0-9]{5,20}\b", text)))

def extract_ids(uploaded):
    name = uploaded.name.lower()
    data = uploaded.getvalue()

    if name.endswith((".xlsx", ".xls")):
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype=str)
        ids = []
        for df in sheets.values():
            ids.extend(extract_ids_from_dataframe(df))
        return list(dict.fromkeys(ids))

    if name.endswith(".csv"):
        df = pd.read_csv(io.BytesIO(data), dtype=str)
        return list(dict.fromkeys(extract_ids_from_dataframe(df)))

    if name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore").upper()
        return list(dict.fromkeys(re.findall(r"\b[A-Z0-9]{5,20}\b", text)))

    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                text += "\n" + "\t".join(cell.text for cell in row.cells)
        return list(dict.fromkeys(re.findall(r"\b[A-Z0-9]{5,20}\b", text.upper())))

    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
        return list(dict.fromkeys(re.findall(r"\b[A-Z0-9]{5,20}\b", text.upper())))

    raise ValueError("Supported files: XLSX, XLS, CSV, PDF, DOCX, TXT.")

def show_matches(ids, con):
    lookup = dict(con.execute("SELECT neo_id,name FROM neo_master").fetchall())
    matched_ids = [i for i in ids if i in lookup]
    result = pd.DataFrame({
        "NEO ID": matched_ids,
        "Name": [lookup[i] for i in matched_ids]
    })
    st.write(f"Found **{len(ids)} unique NEO IDs** and matched **{len(result)}** of them.")
    st.metric("Matches found", len(result))
    st.dataframe(result, use_container_width=True, hide_index=True)
    st.caption("Copy matched results")
    st.code(result.to_csv(index=False, sep="\t"), language=None)

    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        result.to_excel(writer, index=False, sheet_name="Matched")
    st.download_button(
        "Download matched Excel",
        data=out.getvalue(),
        file_name="neo_id_matched_result.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

con = seed_database()
master = get_master(con)

with st.sidebar:
    st.header("Master database")
    st.metric("Stored NEO IDs", len(master))

    master_file = st.file_uploader(
        "Add/replace master list",
        type=["xlsx", "xls", "csv"],
        help="Use a file containing NEO IDs and names. Existing IDs are updated; other IDs are retained."
    )
    if master_file is not None:
        try:
            if master_file.name.lower().endswith(".csv"):
                mdf = pd.read_csv(master_file, dtype=str)
            else:
                sheets = pd.read_excel(master_file, sheet_name=None, dtype=str)
                # Use the sheet containing likely NEO/name columns.
                mdf = next(
                    (d for d in sheets.values()
                     if any("neo" in str(c).lower() and "id" in str(c).lower() for c in d.columns)
                     and any("name" in str(c).lower() for c in d.columns)),
                    next(iter(sheets.values()))
                )
            cols = {str(c).lower(): c for c in mdf.columns}
            neo_col = next((c for c in mdf.columns if "neo" in str(c).lower() and "id" in str(c).lower()), None)
            name_col = next((c for c in mdf.columns if "full name" in str(c).lower() or str(c).lower()=="name" or "name" in str(c).lower()), None)
            if not neo_col or not name_col:
                st.error("Could not identify NEO ID and Name columns.")
            else:
                rows = []
                for _, r in mdf[[neo_col, name_col]].dropna().iterrows():
                    nid, nm = normalize_id(r[neo_col]), str(r[name_col]).strip()
                    if nid and nm:
                        rows.append((nid, nm))
                con.executemany("INSERT OR REPLACE INTO neo_master(neo_id,name) VALUES (?,?)", rows)
                con.commit()
                st.success(f"Added/updated {len(rows)} records.")
                st.rerun()
        except Exception as e:
            st.error(f"Could not import master list: {e}")

paste_text = st.text_area(
    "Paste NEO IDs",
    placeholder="Paste or type one or more NEO IDs, separated by spaces, commas, or new lines.",
    key="pasted_ids"
)
if st.button("Match pasted IDs", disabled=not paste_text.strip()):
    ids = list(dict.fromkeys(re.findall(r"\b[A-Z0-9]{5,20}\b", paste_text.upper())))
    show_matches(ids, con)

uploaded = st.file_uploader(
    "Upload XLSX, XLS, CSV, PDF, DOCX, or TXT",
    type=["xlsx", "xls", "csv", "pdf", "docx", "txt"],
    key="later_file"
)

if uploaded:
    try:
        ids = extract_ids(uploaded)
        show_matches(ids, con)
    except Exception as e:
        st.error(str(e))

with st.expander("View master database"):
    st.dataframe(master, use_container_width=True, hide_index=True)
